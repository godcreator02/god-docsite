"""Shared style/field specification; no equation conversion is shared between routes."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
ROOT = Path(__file__).resolve().parents[1]
FIXTURE = json.loads((ROOT / 'fixtures/canonical.json').read_text(encoding='utf-8'))

def element(tag, **attrs):
    e = OxmlElement(tag)
    for k,v in attrs.items(): e.set(qn(k), str(v))
    return e

def field(parent, code, value='1', bookmark=None):
    if bookmark:
        parent.append(element('w:bookmarkStart', **{'w:id':bookmark[0], 'w:name':bookmark[1]}))
    f=element('w:fldSimple', **{'w:instr':code})
    r=element('w:r'); t=element('w:t'); t.text=value; r.append(t); f.append(r); parent.append(f)
    if bookmark: parent.append(element('w:bookmarkEnd', **{'w:id':bookmark[0]}))

def fonts(style, cn='仿宋', size=12, bold=False):
    style.font.name='Times New Roman'; style.font.size=Pt(size); style.font.bold=bold; style.font.italic=False
    style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(rf.attrib):
        if 'theme' in key.lower(): del rf.attrib[key]
    rf.set(qn('w:eastAsia'),cn)
    for b in list(style.element.findall('.//'+qn('w:pBdr'))): b.getparent().remove(b)

def styled_document():
    d=Document(); s=d.sections[0]
    s.page_width=Mm(210); s.page_height=Mm(297)
    s.top_margin=s.bottom_margin=Mm(25.4); s.left_margin=Mm(23); s.right_margin=Mm(22)
    s.header_distance=s.footer_distance=Mm(12.7); s.different_first_page_header_footer=True
    for name in ['Normal','Body Text','First Paragraph','Compact','Cover Heading','Cover Meta','Equation','TOC Title','TOC 1','TOC 2']:
        if name not in d.styles: d.styles.add_style(name,1)
        st=d.styles[name]; fonts(st)
        st.paragraph_format.space_before=Pt(0); st.paragraph_format.space_after=Pt(0)
        st.paragraph_format.line_spacing=1.5; st.paragraph_format.first_line_indent=Pt(24)
        st.paragraph_format.widow_control=True
    for name,size in [('Heading 1',15),('Heading 2',12)]:
        st=d.styles[name]; fonts(st,size=size,bold=True); p=st.paragraph_format
        p.first_line_indent=Pt(0); p.space_before=Pt(12); p.space_after=Pt(6); p.line_spacing=1.5; p.keep_with_next=True
    for name,size,bold in [('Title',36,True),('Cover Heading',26,False),('Cover Meta',16,False),('TOC Title',18,True)]:
        st=d.styles[name]; fonts(st,size=size,bold=bold); p=st.paragraph_format
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.first_line_indent=Pt(0); p.line_spacing=1.5; p.space_after=Pt(0)
    d.styles['Cover Heading'].paragraph_format.space_before=Pt(70)
    d.styles['Title'].paragraph_format.space_before=Pt(10)
    d.styles['Cover Meta'].paragraph_format.space_before=Pt(18)
    for name in ['Equation','Caption','TOC 1','TOC 2']:
        st=d.styles[name]; fonts(st); st.paragraph_format.first_line_indent=Pt(0)
        st.paragraph_format.line_spacing=1.5
    eq=d.styles['Equation'].paragraph_format
    eq.space_before=Pt(6); eq.space_after=Pt(6)
    eq.tab_stops.add_tab_stop(Mm(82.5), WD_TAB_ALIGNMENT.CENTER)
    eq.tab_stops.add_tab_stop(Mm(165), WD_TAB_ALIGNMENT.RIGHT)
    eq.keep_together=True
    d.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.styles['Caption'].paragraph_format.space_before=d.styles['Caption'].paragraph_format.space_after=Pt(0)
    d.styles['Caption'].paragraph_format.keep_with_next=True
    for name in ['Header','Footer']:
        fonts(d.styles[name],size=10.5)
        d.styles[name].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    s.header.paragraphs[0].text=FIXTURE['layout']['headerText']
    field(s.footer.paragraphs[0]._p,' PAGE ')
    d.settings.element.append(element('w:updateFields',**{'w:val':'true'}))
    return d

def cover(d):
    d.add_paragraph('国网科技报告风格样稿','Cover Heading')
    d.add_paragraph('技术路线对比报告','Title')
    p=d.add_paragraph('项目名称  '+FIXTURE['title'],'Cover Meta'); p.paragraph_format.space_before=Pt(100)
    d.add_paragraph('文档用途  排版与公式转换实验','Cover Meta')
    d.add_paragraph('编制日期  2026 年 10 月 10 日','Cover Meta')
    d.add_page_break(); d.add_paragraph('目录','TOC Title')
    p=d.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0)
    field(p._p,' TOC \\o "1-2" \\h \\z \\u ','目录将在 Word 中更新')
    d.add_page_break()

def parameter_table(d):
    p=d.add_paragraph(style='Caption'); p.add_run('表 '); field(p._p,' SEQ Table \\* ARABIC ',bookmark=('20','TableParameters')); p.add_run(' '+FIXTURE['table']['caption'])
    t=d.add_table(rows=1,cols=4); t.style='Table Grid'; t.autofit=False
    widths=FIXTURE['table']['columnWidthsMm']
    for col,w in zip(t.columns,widths): col.width=Mm(w)
    data=[FIXTURE['table']['headers']]
    for i in range(40):
        row=list(FIXTURE['table']['rows'][i%4]); row[0]+=f' {i+1}'; data.append(row)
    for i,row in enumerate(data):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        tr=t.rows[i]._tr.get_or_add_trPr(); tr.append(element('w:cantSplit'))
        if i==0: tr.append(element('w:tblHeader'))
        for j,(c,txt) in enumerate(zip(cells,row)):
            c.width=Mm(widths[j]); c.text=txt
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Pt(0); p.paragraph_format.line_spacing=1.5
                p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs: r.bold=(i==0)
    return t
